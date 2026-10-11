"""Apresentação PowerPoint do painel executivo (.pptx): gráficos nativos, um slide por loja, transições e animações."""
import io
import zipfile

import pytest
from pptx import Presentation

from app.models import AuditLog
from tests.test_director import _user, _world

URL = "/diretoria/export.pptx?start=2026-01&end=2026-02"


def _deck(client):
    r = client.get(URL)
    assert r.status_code == 200, r.status_code
    return r, Presentation(io.BytesIO(r.content))


def test_director_downloads_a_valid_pptx_with_native_charts(client, db):
    _world(db)
    r, prs = _deck(_user(client, "diretor_ppt", "director"))
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.presentationml.presentation")
    assert 'filename="painel-diretoria-economart.pptx"' in r.headers["content-disposition"]
    charts = [s for sl in prs.slides for s in sl.shapes if s.has_chart]
    assert len(charts) >= 4 + 3                                  # 4 do painel + 1 por loja (3 lojas no cenário)
    assert prs.slide_width > prs.slide_height                     # 16:9


def test_deck_has_one_slide_per_store_and_the_expected_structure(client, db):
    _world(db)
    _, prs = _deck(_user(client, "diretor_ppt2", "director"))
    titles = [" ".join(sh.text_frame.text for sh in sl.shapes if sh.has_text_frame and sh.name == "Título") for sl in prs.slides]
    assert titles[0].startswith("Controle de Energia") and titles[-1].startswith("Dúvidas")
    assert any("Loja A" in t for t in titles) and any("Loja B" in t for t in titles) and any("Loja C" in t for t in titles)
    assert any(t.startswith("Destaques") for t in titles) and any("Ranking" in t for t in titles)


def test_every_slide_has_a_transition_and_an_entrance_animation(client, db):
    _world(db)
    r = _user(client, "diretor_ppt3", "director").get(URL)
    z = zipfile.ZipFile(io.BytesIO(r.content))
    slides = [n for n in z.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml")]
    assert len(slides) >= 10
    for n in slides:
        xml = z.read(n).decode("utf-8")
        assert "<p:transition" in xml and "<p:timing" in xml, n
        assert xml.index("<p:transition") < xml.index("<p:timing")          # ordem do esquema OOXML


def test_deck_carries_only_the_dashboard_numbers(client, db):
    _world(db)
    c = _user(client, "diretor_ppt4", "director")
    _, prs = _deck(c)
    text = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
    assert "JAN/2026 a FEV/2026" in text.replace("  ", " ") or "JAN/2026" in text
    assert "Gerado por diretor_ppt4" in text


def test_pptx_is_for_director_and_admin_only_and_is_audited(client, db):
    _world(db)
    assert _user(client, "func_ppt", "operator").get(URL).status_code == 403
    assert _user(client, "cons_ppt", "viewer").get(URL).status_code == 403
    assert client.get(URL).status_code == 200                                   # administrador
    assert any(a.action == "export" and "pptx" in str(a.details) for a in db.query(AuditLog).all())


def test_pptx_with_no_data_still_opens(client):
    r = _user(client, "diretor_vazio", "director").get(URL)
    assert r.status_code == 200
    prs = Presentation(io.BytesIO(r.content))
    assert len(prs.slides) >= 3


def test_dashboard_has_the_pptx_button_and_the_dependency_is_declared(client, db):
    _world(db)
    html = _user(client, "diretor_ppt5", "director").get("/diretoria?start=2026-01&end=2026-02").text
    assert "/diretoria/export.pptx" in html and "Baixar apresentação" in html
    assert "python-pptx" in open("requirements.txt", encoding="utf-8").read()


@pytest.mark.parametrize("topic", ["dir-exportar", "apresentacao-sistema"])
def test_help_covers_the_new_topics(client, db, topic):
    _world(db)
    page = _user(client, f"diretor_h_{topic[:3]}", "director").get("/ajuda").text
    assert f'id="{topic}"' in page
