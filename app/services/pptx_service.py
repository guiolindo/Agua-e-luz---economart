"""Apresentação (PowerPoint .pptx) do painel executivo, no mesmo visual do modo apresentação do sistema.

Para reuniões em que o diretor ou o gerente precisa de um arquivo PowerPoint: capa, destaques, gráficos do painel e um
slide por loja com a evolução mensal. Os gráficos são NATIVOS do PowerPoint (editáveis: clique com o botão direito >
Editar dados), as cores e a ordem das lojas seguem o painel, e cada slide tem transição suave e entrada em cascata dos
elementos (animações simples de esmaecer, sem depender de nada instalado).

Só usa dados que o painel já calcula (executive_service.build); nenhum número novo é inventado aqui.
"""
import io
import itertools
from datetime import datetime

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from app.utils import formatting as fmt

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
MAX_STORE_SLIDES = 60

# paleta do modo apresentação
NAVY, NAVY_2, CARD = "091D38", "0D2A50", "12315A"
ORANGE, AMBER, SKY = "FF7A1A", "FFB15C", "4AA3FF"
WHITE, MUTED, GRID = "FFFFFF", "AEBBD0", "27456F"
GOOD, BAD = "3DBE7A", "FF6B5E"
HEAD_FONT, BODY_FONT = "Segoe UI Semibold", "Segoe UI"

W, H = 13.333, 7.5
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def _rgb(hex6: str) -> RGBColor:
    return RGBColor.from_string(hex6.upper())


# ---------------------------------------------------------------- animações (transição + entrada em cascata)
def _transition(slide, kind: str = "fade") -> None:
    """Transição suave entre slides (esmaecer). Vai depois de <p:clrMapOvr> e antes de <p:timing> (ordem do esquema)."""
    sld = slide._element
    xml = f'<p:transition xmlns:p="{NS_P}" spd="med"><p:{kind}/></p:transition>'
    sld.append(etree.fromstring(xml))


def _fade_in(slide, steps: list[list[int]], dur: int = 450, gap: int = 120) -> None:
    """Entrada em cascata (esmaecer) sem clique: cada passo é uma lista de formas que aparecem juntas; os passos vão
    um depois do outro, começando sozinhos quando o slide abre."""
    if not steps:
        return
    ids = itertools.count(4)
    pars, t = [], 0
    for step in steps:
        effects = []
        for k, spid in enumerate(step):
            a, b, c = next(ids), next(ids), next(ids)
            node = "afterEffect" if k == 0 else "withEffect"
            effects.append(
                f'<p:par><p:cTn id="{a}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" nodeType="{node}">'
                f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                f'<p:set><p:cBhvr><p:cTn id="{b}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>'
                f'</p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
                f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{c}" dur="{dur}"/>'
                f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect>'
                f'</p:childTnLst></p:cTn></p:par>')
        pars.append(f'<p:par><p:cTn id="{next(ids)}" fill="hold"><p:stCondLst><p:cond delay="{t}"/></p:stCondLst>'
                    f'<p:childTnLst>{"".join(effects)}</p:childTnLst></p:cTn></p:par>')
        t += dur + gap
    xml = (f'<p:timing xmlns:p="{NS_P}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
           f'<p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
           f'<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>{"".join(pars)}</p:childTnLst></p:cTn></p:par>'
           f'</p:childTnLst></p:cTn>'
           f'<p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>'
           f'</p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>')
    slide._element.append(etree.fromstring(xml))


# ---------------------------------------------------------------- peças visuais
def _new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])         # layout "Em branco"
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = _rgb(NAVY)
    return slide


def _text(slide, x, y, w, h, text, size=14, bold=False, color=WHITE, align=PP_ALIGN.LEFT, font=BODY_FONT,
          anchor=MSO_ANCHOR.TOP, name=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
        r.font.color.rgb = _rgb(color)
    return box


def _kicker_title(slide, kicker, title, sub=None):
    k = _text(slide, 0.7, 0.55, 11.5, 0.3, kicker.upper(), size=12, bold=True, color=AMBER, name="Chamada")
    t = _text(slide, 0.7, 0.9, 11.9, 0.9, title, size=34, bold=True, color=WHITE, font=HEAD_FONT, name="Título")
    out = [k, t]
    if sub:
        out.append(_text(slide, 0.7, 1.78, 11.9, 0.4, sub, size=14, color=MUTED, name="Subtítulo"))
    return out


def _card(slide, x, y, w, h, value, label, value_color=WHITE, note=None, name="Indicador"):
    """Cartão de indicador: UMA forma com valor grande + rótulo (+ nota), para animar de uma vez."""
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.name = name
    shp.adjustments[0] = 0.08
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(CARD)
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.22)
    tf.margin_top = tf.margin_bottom = Inches(0.15)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    vsize = 28 if len(value) <= 8 else 22 if len(value) <= 12 else 18
    for i, (txt, size, bold, color, font) in enumerate((
            (label.upper(), 11, True, MUTED, BODY_FONT), (value, vsize, True, value_color, HEAD_FONT), (note or "", 11, False, MUTED, BODY_FONT))):
        if i == 2 and not note:
            continue
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = txt
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
        r.font.color.rgb = _rgb(color)
    return shp


def _style_chart(chart, legend=False):
    chart.font.size = Pt(12)
    chart.font.name = BODY_FONT
    chart.font.color.rgb = _rgb(MUTED)
    chart.has_title = False
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(12)
        chart.legend.font.color.rgb = _rgb(MUTED)
    cat, val = chart.category_axis, chart.value_axis
    cat.format.line.color.rgb = _rgb(GRID)
    cat.has_major_gridlines = False
    cat.tick_labels.font.size = Pt(12)
    cat.tick_labels.font.color.rgb = _rgb(MUTED)
    val.format.line.fill.background()
    val.has_major_gridlines = True
    val.major_gridlines.format.line.color.rgb = _rgb(GRID)
    val.major_gridlines.format.line.width = Pt(0.5)
    val.tick_labels.font.size = Pt(11)
    val.tick_labels.font.color.rgb = _rgb(MUTED)
    val.tick_labels.number_format = '"R$ "#,##0'
    val.tick_labels.number_format_is_linked = False


def _color_points(series, colors):
    for i, c in enumerate(colors):
        pt = series.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = _rgb(c)


def _short(code, name, limit=26):
    label = name if name and str(code).isdigit() else (f"{code} · {name}" if name else code)   # loja numerada: o nome
    return label if len(label) <= limit else label[: limit - 1] + "…"


def _brl(v, dec=0):
    return fmt.brl(v, dec) if v is not None else "—"


def _pct(v):
    return "—" if v is None else f"{'↑' if v > 0 else '↓' if v < 0 else ''} {fmt.pct(abs(v))}".strip()


def _notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ---------------------------------------------------------------- slides
def _cover(prs, title, subtitle, kicker, label):
    s = _new_slide(prs)
    heights = [1.2, 1.9, 1.5, 2.6, 2.0, 3.3, 2.5, 3.9, 3.0, 4.6, 3.6, 5.2]       # barras de energia, como na capa do sistema
    bars = []
    for i, hgt in enumerate(heights):
        shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.3 + i * 0.47), Inches(7.5 - hgt), Inches(0.34), Inches(hgt))
        shp.name = f"Barra {i + 1}"
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(SKY if i % 3 == 2 else ORANGE)
        shp.line.fill.background()
        shp.shadow.inherit = False
        bars.append(shp.shape_id)
    k = _text(s, 0.8, 2.0, 6.4, 0.4, kicker.upper(), size=13, bold=True, color=AMBER, name="Chamada")
    t = _text(s, 0.8, 2.5, 6.6, 2.2, title, size=54, bold=True, font=HEAD_FONT, name="Título")
    sub = _text(s, 0.8, 4.9, 6.2, 0.9, subtitle, size=18, color=MUTED, name="Subtítulo")
    lab = _text(s, 0.8, 6.3, 6.2, 0.4, label, size=12, color=MUTED, name="Período")
    _transition(s)
    _fade_in(s, [[k.shape_id], [t.shape_id], [sub.shape_id, lab.shape_id], bars], dur=500)
    return s


def _highlights(prs, data, insights):
    k = data["kpi"]
    s = _new_slide(prs)
    head = _kicker_title(s, "Visão geral", f"Destaques de {data['focus_label'] or 'o período'}")
    ids = [[h.shape_id for h in head]]
    tone_color = {"good": GOOD, "bad": BAD, "info": SKY}
    y = 2.05
    bullets = []
    for ins in insights[:6]:
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.72), Inches(y + 0.09), Inches(0.16), Inches(0.16))
        dot.name = "Marcador"
        dot.fill.solid()
        dot.fill.fore_color.rgb = _rgb(tone_color.get(ins["tone"], SKY))
        dot.line.fill.background()
        dot.shadow.inherit = False
        txt = _text(s, 1.05, y, 5.9, 0.78, ins["text"], size=14, color=WHITE, name="Destaque")
        bullets.append([dot.shape_id, txt.shape_id])
        y += 0.82
    cards = [
        _card(s, 7.35, 1.95, 2.7, 1.55, _brl(k["total"]), "Total no período", note=f"{len(data['months'])} meses · {k['stores']} lojas"),
        _card(s, 10.2, 1.95, 2.7, 1.55, _brl(k["last_month"]), f"Último mês ({data['focus_label']})" if data["focus_label"] else "Último mês",
              note=_pct(k["last_vs_prev"]["pct"]) + (f" vs. {data['prev_label']}" if data["prev_label"] else "") if k["last_vs_prev"] else None),
        _card(s, 7.35, 3.7, 2.7, 1.55, _brl(k["rs_kwh"], 3) if k["rs_kwh"] else "—", "Custo efetivo por kWh", note="reais por kWh: total da conta ÷ kWh"),
        _card(s, 10.2, 3.7, 2.7, 1.55, str(k["pending"]), f"Pendências de {data['focus_label']}" if data["focus_label"] else "Pendências",
              value_color=AMBER if k["pending"] else GOOD, note="contas ou lançamentos do mês"),
    ]
    if k.get("top"):
        top = k["top"]
        cards.append(_card(s, 7.35, 5.45, 5.55, 1.2, _short(top["code"], top["name"], 34), "Maior gasto no período", value_color=ORANGE,
                           note=f"{_brl(top['total'])} · {fmt.pct(top['share'])} do total"))
    ids += bullets + [[c.shape_id] for c in cards]
    _transition(s)
    _fade_in(s, ids)
    _notes(s, " ".join(i["text"] for i in insights[:6]))
    return s


def _monthly_by_supplier(prs, data):
    s = _new_slide(prs)
    note = ("* " + "; ".join(data["partial_labels"]) + " — meses com poucas lojas lançadas.") if data["partial_labels"] else None
    head = _kicker_title(s, "Evolução da empresa", "Gasto por mês e por fornecedor", note)
    cd = CategoryChartData()
    cd.categories = [lbl + ("*" if p else "") for lbl, p in zip(data["labels"], data["partial"])]
    for t in data["by_type"]:
        cd.add_series(t["name"], [v if v else None for v in t["values"]])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(0.6), Inches(2.3), Inches(12.1), Inches(4.9), cd)
    gf.name = "Gráfico: gasto por mês e fornecedor"
    ch = gf.chart
    _style_chart(ch, legend=True)
    ch.plots[0].gap_width = 45
    for ser, t in zip(ch.plots[0].series, data["by_type"]):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = _rgb(t["color"].lstrip("#"))
    _transition(s)
    _fade_in(s, [[h.shape_id for h in head], [gf.shape_id]])
    return s


def _ranking(prs, data):
    rows = data["rows"][:15]
    s = _new_slide(prs)
    head = _kicker_title(s, "Desempenho das lojas", f"Ranking de lojas por gasto (Top {len(rows)})", "total no período, em R$")
    cd = CategoryChartData()
    cd.categories = [_short(r["code"], r["name"], 30) for r in rows]
    cd.add_series("Total no período", [r["total"] for r in rows])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(2.3), Inches(12.1), Inches(4.95), cd)
    gf.name = "Gráfico: ranking de lojas"
    ch = gf.chart
    _style_chart(ch)
    ch.category_axis.reverse_order = True
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.plots[0].gap_width = 40
    ch.plots[0].has_data_labels = True
    dl = ch.plots[0].data_labels
    dl.number_format, dl.number_format_is_linked = '"R$ "#,##0', False
    dl.font.size, dl.font.color.rgb = Pt(11), _rgb(WHITE)
    _color_points(ch.plots[0].series[0], [ORANGE if i == 0 else SKY for i in range(len(rows))])
    _transition(s)
    _fade_in(s, [[h.shape_id for h in head], [gf.shape_id]])
    return s


def _movers(prs, data):
    movers = [r for r in data["rows"] if r["vs_last_month"]["pct"] is not None]
    if not movers:
        return None
    ups = sorted((r for r in movers if r["vs_last_month"]["pct"] > 0), key=lambda r: -r["vs_last_month"]["pct"])[:7]
    downs = sorted((r for r in movers if r["vs_last_month"]["pct"] < 0), key=lambda r: r["vs_last_month"]["pct"])[:7]
    items = ups + downs
    s = _new_slide(prs)
    head = _kicker_title(s, "Desempenho das lojas", f"Maiores altas e quedas em {data['focus_label']}",
                         f"variação contra {data['prev_label']}: vermelho gastou mais, verde gastou menos")
    cd = CategoryChartData()
    cd.categories = [_short(r["code"], r["name"], 28) for r in items]
    cd.add_series("Variação (%)", [round(r["vs_last_month"]["pct"], 2) for r in items])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(2.3), Inches(12.1), Inches(4.95), cd)
    gf.name = "Gráfico: maiores altas e quedas"
    ch = gf.chart
    _style_chart(ch)
    ch.category_axis.reverse_order = True
    ch.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.plots[0].gap_width = 40
    ch.plots[0].has_data_labels = True
    dl = ch.plots[0].data_labels
    dl.number_format, dl.number_format_is_linked = '0.0"%"', False
    dl.font.size, dl.font.color.rgb = Pt(11), _rgb(WHITE)
    vals = [r["vs_last_month"]["pct"] for r in items]
    ch.value_axis.minimum_scale = min(0, min(vals)) * 1.35
    ch.value_axis.maximum_scale = max(0, max(vals)) * 1.35 + 1
    _color_points(ch.plots[0].series[0], [BAD if r["vs_last_month"]["pct"] > 0 else GOOD for r in items])
    _transition(s)
    _fade_in(s, [[h.shape_id for h in head], [gf.shape_id]])
    return s


def _kwh(prs, data):
    rows = sorted((r for r in data["rows"] if r["rs_kwh"]), key=lambda r: -r["rs_kwh"])[:12]
    if not rows:
        return None
    avg = data["kpi"]["rs_kwh"]
    s = _new_slide(prs)
    head = _kicker_title(s, "Custos", "Custo efetivo por kWh", (f"média da empresa {_brl(avg, 3)} por kWh · " if avg else "") + "total da conta ÷ kWh")
    cd = CategoryChartData()
    cd.categories = [_short(r["code"], r["name"], 28) for r in rows]
    cd.add_series("R$ por kWh", [round(r["rs_kwh"], 4) for r in rows])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(2.3), Inches(12.1), Inches(4.95), cd)
    gf.name = "Gráfico: custo por kWh"
    ch = gf.chart
    _style_chart(ch)
    ch.category_axis.reverse_order = True
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.plots[0].gap_width = 40
    ch.plots[0].has_data_labels = True
    dl = ch.plots[0].data_labels
    dl.number_format, dl.number_format_is_linked = '"R$ "0.000', False
    dl.font.size, dl.font.color.rgb = Pt(11), _rgb(WHITE)
    _color_points(ch.plots[0].series[0], [ORANGE if avg and r["rs_kwh"] > avg else SKY for r in rows])
    _transition(s)
    _fade_in(s, [[h.shape_id for h in head], [gf.shape_id]])
    return s


def _divider(prs, kicker, title, sub):
    s = _new_slide(prs)
    k = _text(s, 0.8, 2.6, 11, 0.4, kicker.upper(), size=13, bold=True, color=AMBER, name="Chamada")
    t = _text(s, 0.8, 3.1, 11.5, 1.4, title, size=48, bold=True, font=HEAD_FONT, name="Título")
    sb = _text(s, 0.8, 4.7, 11, 0.8, sub, size=18, color=MUTED, name="Subtítulo")
    _transition(s)
    _fade_in(s, [[k.shape_id], [t.shape_id], [sb.shape_id]])
    return s


def _store_slide(prs, data, r, position, total_stores):
    s = _new_slide(prs)
    sub = " · ".join(x for x in (r.get("region"), f"{position}º de {total_stores} em gasto no período") if x)
    head = _kicker_title(s, "Evolução da loja", _short(r["code"], r["name"], 44), sub)
    lv = r["vs_last_month"]
    cards = [
        _card(s, 0.7, 2.35, 2.9, 1.35, _brl(r["total"]), "Total no período", note=f"{fmt.pct(r['share'])} do total da empresa"),
        _card(s, 3.75, 2.35, 2.9, 1.35, _brl(r["last"]), f"Último mês ({data['focus_label']})" if data["focus_label"] else "Último mês",
              value_color=BAD if (lv["pct"] or 0) > 0 else GOOD if (lv["pct"] or 0) < 0 else WHITE,
              note=(_pct(lv["pct"]) + (f" vs. {data['prev_label']}" if data["prev_label"] else "")) if lv["pct"] is not None else None),
        _card(s, 6.8, 2.35, 2.9, 1.35, _brl(r["rs_kwh"], 3) if r["rs_kwh"] else "—", "Custo por kWh (R$)",
              note=f"média da empresa {_brl(data['kpi']['rs_kwh'], 3)}" if data["kpi"]["rs_kwh"] else None),
        _card(s, 9.85, 2.35, 2.9, 1.35, f"{fmt.num(r['demand_use'], 0)}%" if r["demand_use"] is not None else "—", "Demanda usada",
              value_color=BAD if (r["demand_over"] or 0) else WHITE,
              note=(f"{r['demand_over']} conta(s) acima do contratado" if r["demand_over"] else "da demanda contratada"))
        if data.get("has_demand") else
        _card(s, 9.85, 2.35, 2.9, 1.35, str(r.get("n_months") or 0), "Meses com conta", note="no período"),
    ]
    cd = CategoryChartData()
    cd.categories = [lbl + ("*" if p else "") for lbl, p in zip(data["labels"], data["partial"])]
    cd.add_series("Gasto da loja", [v if v else None for v in r["values"]])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(3.95), Inches(12.1), Inches(3.3), cd)
    gf.name = "Gráfico: evolução mensal da loja"
    ch = gf.chart
    _style_chart(ch)
    ch.plots[0].gap_width = 45
    ch.plots[0].has_data_labels = True
    dl = ch.plots[0].data_labels
    dl.number_format, dl.number_format_is_linked = '#,##0', False
    dl.font.size, dl.font.color.rgb = Pt(10), _rgb(WHITE)
    last_i = max((i for i, v in enumerate(r["values"]) if v), default=None)
    _color_points(ch.plots[0].series[0], [ORANGE if i == last_i else SKY for i in range(len(r["values"]))])
    _transition(s)
    _fade_in(s, [[h.shape_id for h in head]] + [[c.shape_id] for c in cards] + [[gf.shape_id]], dur=400, gap=90)
    return s


def _pending(prs, data, pending):
    s = _new_slide(prs)
    head = _kicker_title(s, "Qualidade dos dados", f"Pendências de {data['focus_label']}" if data["focus_label"] else "Pendências",
                         "contas ou lançamentos que ainda faltam" if pending else None)
    ids = [[h.shape_id for h in head]]
    if not pending:
        t = _text(s, 0.7, 2.6, 11, 1.0, "Nenhuma pendência no mês. Todas as lojas estão em dia.", size=22, color=GOOD)
        ids.append([t.shape_id])
    else:
        rows = pending[:10]
        gf = s.shapes.add_table(len(rows) + 1, 3, Inches(0.7), Inches(2.4), Inches(11.9), Inches(0.5 * (len(rows) + 1)))
        gf.name = "Tabela de pendências"
        tbl = gf.table
        tbl._tbl.tblPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId").text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
        for col, wid in enumerate((2.6, 4.0, 5.3)):
            tbl.columns[col].width = Inches(wid)
        for ci, head_txt in enumerate(("Loja", "Fornecedor", "Situação")):
            _cell(tbl.cell(0, ci), head_txt, bold=True, color=AMBER, fill=NAVY_2)
        for ri, p in enumerate(rows, 1):
            store = p["store"]
            _cell(tbl.cell(ri, 0), store.short_label if str(store.code).isdigit() else f"{store.code}" + (f" · {store.name}" if store.name else ""), fill=CARD)
            _cell(tbl.cell(ri, 1), p["type"].name, fill=CARD)
            _cell(tbl.cell(ri, 2), p["text"], color=MUTED, fill=CARD)
        ids.append([gf.shape_id])
        if len(pending) > 10:
            more = _text(s, 0.7, 2.5 + 0.5 * (len(rows) + 1) + 0.1, 11, 0.4, f"+ {len(pending) - 10} pendências no painel.", size=12, color=MUTED)
            ids.append([more.shape_id])
    _transition(s)
    _fade_in(s, ids)
    return s


def _cell(cell, text, bold=False, color=WHITE, fill=CARD):
    cell.fill.solid()
    cell.fill.fore_color.rgb = _rgb(fill)
    cell.margin_left = cell.margin_right = Inches(0.15)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.name = Pt(13), bold, BODY_FONT
    r.font.color.rgb = _rgb(color)


def _closing(prs, who, when):
    s = _new_slide(prs)
    heights = [5.2, 3.6, 4.6, 3.0, 3.9, 2.5, 3.3, 2.0, 2.6, 1.5, 1.9, 1.2]
    bars = []
    for i, hgt in enumerate(heights):
        shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7.3 + i * 0.47), Inches(7.5 - hgt), Inches(0.34), Inches(hgt))
        shp.name = f"Barra {i + 1}"
        shp.fill.solid()
        shp.fill.fore_color.rgb = _rgb(SKY if i % 3 == 0 else ORANGE)
        shp.line.fill.background()
        shp.shadow.inherit = False
        bars.append(shp.shape_id)
    k = _text(s, 0.8, 2.6, 6, 0.4, "OBRIGADO", size=13, bold=True, color=AMBER, name="Chamada")
    t = _text(s, 0.8, 3.1, 6.5, 1.4, "Dúvidas?", size=60, bold=True, font=HEAD_FONT, name="Título")
    sb = _text(s, 0.8, 4.9, 6.2, 1.2, ["Dados do sistema Controle de Energia.",
                                       f"Gerado por {who} em {when}."], size=14, color=MUTED, name="Rodapé")
    _transition(s)
    _fade_in(s, [[k.shape_id], [t.shape_id], [sb.shape_id], bars], dur=500)
    return s


# ---------------------------------------------------------------- deck
def executive_deck(data: dict, insights: list[dict], pending: list, filters_label: str, who: str = "o sistema") -> bytes:
    """Monta a apresentação a partir dos dados do painel (`executive_service.build`)."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.title = "Controle de Energia: visão executiva"
    prs.core_properties.author = "Economart · Controle de Energia"
    prs.core_properties.subject = filters_label
    prs.slide_master.background.fill.solid()
    prs.slide_master.background.fill.fore_color.rgb = _rgb(NAVY)

    period = (f"{fmt.month_label(data['start'])} a {fmt.month_label(data['end'])}" if data["months"] else "sem dados no período")
    _cover(prs, "Controle de Energia", "Visão executiva", "Economart Atacadista", f"{period} · {filters_label}")
    rows = data["rows"]
    if not rows:
        s = _new_slide(prs)
        t = _text(s, 0.8, 3.0, 11.5, 1.2, "Ainda não há contas lançadas neste período.", size=28, color=WHITE, font=HEAD_FONT)
        _transition(s)
        _fade_in(s, [[t.shape_id]])
    else:
        _highlights(prs, data, insights)
        _monthly_by_supplier(prs, data)
        _ranking(prs, data)
        _movers(prs, data)
        _kwh(prs, data)
        _pending(prs, data, pending)
        shown = rows[:MAX_STORE_SLIDES]
        extra = f" (as {MAX_STORE_SLIDES} maiores de {len(rows)})" if len(rows) > MAX_STORE_SLIDES else ""
        _divider(prs, "Loja a loja", "A evolução de cada loja", f"{len(shown)} lojas, da maior para a menor em gasto no período{extra}")
        for pos, r in enumerate(shown, 1):
            _store_slide(prs, data, r, pos, len(rows))
    _closing(prs, who, datetime.now().strftime("%d/%m/%Y"))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()

