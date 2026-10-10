"""Planilhas XLSX formatadas (painel da diretoria e lista de contas).

Valores ficam como números reais (moeda, percentual), não como texto, para o Excel somar, ordenar e filtrar.
Textos digitados por pessoas nunca viram fórmula: células que começam com = + - @ são gravadas como texto.
"""
import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.utils import formatting as fmt

NAVY, BLUE, ORANGE, SOFT, LINE, GREY = "10243F", "1B4F8A", "F47920", "EAF2FA", "C4CCD6", "5B6573"
GOOD, BAD, WARN_BG = "1D7A46", "B3261E", "FFF3DC"

BRL = '"R$" #,##0.00'
BRL0 = '"R$" #,##0'
PCT = '0.0%;[Red]-0.0%'
KWH = '"R$" 0.000'
EXCEL_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_thin = Side(style="thin", color=LINE)
_BORDER = Border(bottom=_thin)
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def _text(cell, value) -> None:
    """Grava texto sem nunca interpretá-lo como fórmula."""
    cell.value = value
    if isinstance(value, str) and value[:1] in _FORMULA_START:
        cell.data_type = "s"


def _fraction(pct) -> float | None:
    return None if pct is None else float(pct) / 100.0


def _title(ws: Worksheet, title: str, subtitle: str, ncols: int) -> int:
    """Faixa de título no topo (marca). Devolve a próxima linha livre."""
    last = get_column_letter(max(ncols, 4))
    for r, text, size, color, bold in ((1, title, 16, "FFFFFF", True), (2, subtitle, 10, "D6E4F5", False)):
        ws.merge_cells(f"A{r}:{last}{r}")
        c = ws[f"A{r}"]
        c.value = text
        c.font = Font(name="Calibri", size=size, bold=bold, color=color)
        c.alignment = Alignment(vertical="center", indent=1)
        ws.row_dimensions[r].height = 30 if r == 1 else 20
        for col in range(1, max(ncols, 4) + 1):
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=NAVY)
    ws.merge_cells(f"A3:{last}3")
    ws["A3"].fill = PatternFill("solid", fgColor=ORANGE)
    ws.row_dimensions[3].height = 4
    return 5


def _header(ws: Worksheet, row: int, labels: list[str]) -> None:
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.alignment = Alignment(horizontal="left" if i == 1 else "center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 32


def _body_style(ws: Worksheet, first: int, last: int, ncols: int) -> None:
    for r in range(first, last + 1):
        for col in range(1, ncols + 1):
            c = ws.cell(row=r, column=col)
            c.border = _BORDER
            c.font = Font(name="Calibri", size=10, color="1F2937")
            if (r - first) % 2 == 1:
                c.fill = PatternFill("solid", fgColor="F6F8FB")


def _widths(ws: Worksheet, widths: list[float]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _finish(ws: Worksheet, header_row: int, last_row: int, ncols: int, freeze_col: str = "B", landscape: bool = True) -> None:
    ws.freeze_panes = f"{freeze_col}{header_row + 1}"
    if last_row > header_row:
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(ncols)}{last_row}"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.oddFooter.left.text = "Economart · Controle de Energia"
    ws.oddFooter.right.text = "Página &P de &N"


def _variation_colors(ws: Worksheet, ref: str) -> None:
    ws.conditional_formatting.add(ref, CellIsRule(operator="greaterThan", formula=["0"], font=Font(color=BAD, bold=True)))
    ws.conditional_formatting.add(ref, CellIsRule(operator="lessThan", formula=["0"], font=Font(color=GOOD, bold=True)))


def _stamp() -> str:
    from app.models.mixins import utcnow
    from app.utils.timezone import to_local

    return f"Emitido em {to_local(utcnow()).strftime('%d/%m/%Y %H:%M')}"


def _bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- painel da diretoria
def executive_workbook(data: dict, insights: list[dict], pending: list[dict], filters_text: str) -> bytes:
    wb = Workbook()
    focus, prev = data["focus_label"], data["prev_label"]
    period = f"{fmt.month_label(data['start'])} a {fmt.month_label(data['end'])}"
    subtitle = f"{period} · {filters_text} · {_stamp()}"
    k = data["kpi"]

    # ---- Resumo
    ws = wb.active
    ws.title = "Resumo"
    _title(ws, "Painel da diretoria · Custos de energia", subtitle, 4)
    r = 5
    ws.cell(row=r, column=1, value="Indicadores").font = Font(bold=True, size=12, color=NAVY)
    r += 1
    rise, fall = k.get("rise"), k.get("fall")

    def lead(row):
        if not row:
            return "—"
        absolute = row.get("abs_last_month")
        if absolute is None:
            extra = ""
        else:
            extra = ("+" if absolute > 0 else "−" if absolute < 0 else "") + fmt.brl(abs(absolute), 0) + " · "
        return f"{row['code']} ({extra}{row['vs_last_month']['text']})"

    items = [
        ("Total no período", k["total"], BRL0),
        (f"Último mês ({focus})", k["last_month"], BRL0),
        ("Média mensal", k["avg_month"], BRL0),
        ("Custo médio (R$/kWh)", k["rs_kwh"], KWH),
        ("Lojas com dados", k["stores"], "0"),
        ("Maior gasto", f"{k['top']['code']} ({fmt.pct(k['top']['share'])} do total)" if k["top"] else "—", None),
        (f"Maior alta em {focus}", lead(rise), None),
        (f"Maior queda em {focus}", lead(fall), None),
        (f"Pendências de {focus}", len(pending), "0"),
    ]
    for label, value, fmt_ in items:
        a = ws.cell(row=r, column=1, value=label)
        a.font = Font(color=GREY, size=10)
        a.alignment = Alignment(vertical="center")
        b = ws.cell(row=r, column=2)
        _text(b, value)
        b.font = Font(bold=True, size=11, color=NAVY)
        b.alignment = Alignment(horizontal="left")
        if fmt_ and isinstance(value, (int, float)):
            b.number_format = fmt_
        for col in (1, 2):
            ws.cell(row=r, column=col).border = _BORDER
        r += 1
    if data.get("partial_labels"):
        r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        c = ws.cell(row=r, column=1, value="Mês incompleto: " + "; ".join(data["partial_labels"])
                    + ". Marcado com * e fora das variações; os indicadores usam " + focus + ".")
        c.fill = PatternFill("solid", fgColor=WARN_BG)
        c.font = Font(size=10, color="6B4300")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[r].height = 32
        r += 1
    r += 1
    ws.cell(row=r, column=1, value=f"Destaques de {focus}").font = Font(bold=True, size=12, color=NAVY)
    r += 1
    marks = {"bad": "▲ Atenção", "good": "▼ Melhora", "info": "● Informação"}
    for item in insights:
        t = ws.cell(row=r, column=1, value=marks.get(item["tone"], "●"))
        t.font = Font(size=10, bold=True, color={"bad": BAD, "good": GOOD}.get(item["tone"], GREY))
        t.alignment = Alignment(vertical="top")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        c = ws.cell(row=r, column=2)
        _text(c, item["text"])
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(size=10)
        ws.row_dimensions[r].height = 30 if len(item["text"]) > 90 else 18
        r += 1
    _widths(ws, [30, 46, 26, 26])
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # ---- Lojas (comparativo)
    ws = wb.create_sheet("Lojas")
    heads = ["Loja", "Nome", "Região", "Total no período", "% da empresa", "Média por mês", "Variação da média vs. período anterior",
             f"{focus}", f"Variação vs. {prev}", "R$/kWh", "Demanda usada", "Contas acima da contratada"]
    first = _title(ws, "Comparativo entre lojas", subtitle, len(heads))
    _header(ws, first, heads)
    row = first + 1
    for s in data["rows"]:
        _text(ws.cell(row=row, column=1), s["code"])
        _text(ws.cell(row=row, column=2), s["name"])
        _text(ws.cell(row=row, column=3), s["region"])
        vals = [s["total"], _fraction(s["share"]), s["avg"], _fraction(s["vs_prev_period"]["pct"]), s["last"],
                _fraction(s["vs_last_month"]["pct"]), s["rs_kwh"], _fraction(s["demand_use"]), s["demand_over"] or 0]
        for col, v in enumerate(vals, start=4):
            ws.cell(row=row, column=col, value=v)
        row += 1
    last = row - 1
    _body_style(ws, first + 1, last, len(heads))
    for r_ in range(first + 1, last + 1):
        ws.cell(row=r_, column=1).font = Font(bold=True, color=BLUE, size=10)
        for col, f in ((4, BRL0), (5, "0.0%"), (6, BRL0), (7, PCT), (8, BRL0), (9, PCT), (10, KWH), (11, "0%"), (12, "0")):
            ws.cell(row=r_, column=col).number_format = f
            ws.cell(row=r_, column=col).alignment = Alignment(horizontal="right")
    if last >= first + 1:
        _variation_colors(ws, f"G{first + 1}:G{last}")
        _variation_colors(ws, f"I{first + 1}:I{last}")
        ws.conditional_formatting.add(f"K{first + 1}:K{last}", CellIsRule(operator="greaterThan", formula=["1"], font=Font(color=BAD, bold=True)))
        tr = last + 1
        ws.cell(row=tr, column=1, value="Total").font = Font(bold=True, color="FFFFFF")
        ws.cell(row=tr, column=4, value=f"=SUM(D{first + 1}:D{last})").number_format = BRL0
        ws.cell(row=tr, column=8, value=f"=SUM(H{first + 1}:H{last})").number_format = BRL0
        for col in range(1, len(heads) + 1):
            c = ws.cell(row=tr, column=col)
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.font = Font(bold=True, color="FFFFFF", size=10)
    _widths(ws, [12, 28, 9, 17, 12, 15, 20, 15, 15, 11, 12, 16])
    _finish(ws, first, last, len(heads), "B")

    # ---- Mês a mês
    ws = wb.create_sheet("Mês a mês")
    labels = [fmt.month_short(m) + ("*" if data["partial"][i] else "") for i, m in enumerate(data["months"])]
    heads = ["Loja", "Região"] + labels + ["Total"]
    first = _title(ws, "Gasto mensal por loja (R$)", subtitle + " · * mês incompleto", len(heads))
    _header(ws, first, heads)
    row = first + 1
    nm = len(labels)
    for s in data["rows"]:
        _text(ws.cell(row=row, column=1), s["code"])
        _text(ws.cell(row=row, column=2), s["region"])
        for i, v in enumerate(s["values"]):
            ws.cell(row=row, column=3 + i, value=v)
        ws.cell(row=row, column=3 + nm, value=f"=SUM({get_column_letter(3)}{row}:{get_column_letter(2 + nm)}{row})")
        row += 1
    last = row - 1
    _body_style(ws, first + 1, last, len(heads))
    for r_ in range(first + 1, last + 1):
        ws.cell(row=r_, column=1).font = Font(bold=True, color=BLUE, size=10)
        for col in range(3, 4 + nm):
            ws.cell(row=r_, column=col).number_format = BRL0
            ws.cell(row=r_, column=col).alignment = Alignment(horizontal="right")
        ws.cell(row=r_, column=3 + nm).font = Font(bold=True, size=10)
    if last >= first + 1:
        rng = f"C{first + 1}:{get_column_letter(2 + nm)}{last}"
        ws.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="7FA9D6"))
        tr = last + 1
        ws.cell(row=tr, column=1, value="Total")
        for col in range(3, 4 + nm):
            L = get_column_letter(col)
            ws.cell(row=tr, column=col, value=f"=SUM({L}{first + 1}:{L}{last})").number_format = BRL0
        for col in range(1, len(heads) + 1):
            c = ws.cell(row=tr, column=col)
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.font = Font(bold=True, color="FFFFFF", size=10)
    _widths(ws, [12, 9] + [13] * nm + [16])
    _finish(ws, first, last, len(heads), "C")

    # ---- Fornecedores (tipos de registro)
    ws = wb.create_sheet("Fornecedores")
    tmap = {t.id: t.name for t in data["types"]}
    used = [t for t in data["by_type"]]
    heads = ["Loja"] + [t["name"] for t in used] + ["Total"]
    first = _title(ws, "Gasto por fornecedor e tipo de despesa (R$)", subtitle, len(heads))
    _header(ws, first, heads)
    row = first + 1
    for s in data["rows"]:
        _text(ws.cell(row=row, column=1), s["code"])
        for i, t in enumerate(used):
            ws.cell(row=row, column=2 + i, value=s["mix"].get(str(t["id"])) or None)
        ws.cell(row=row, column=2 + len(used), value=f"=SUM(B{row}:{get_column_letter(1 + len(used))}{row})")
        row += 1
    last = row - 1
    _body_style(ws, first + 1, last, len(heads))
    for r_ in range(first + 1, last + 1):
        ws.cell(row=r_, column=1).font = Font(bold=True, color=BLUE, size=10)
        for col in range(2, len(heads) + 1):
            ws.cell(row=r_, column=col).number_format = BRL0
            ws.cell(row=r_, column=col).alignment = Alignment(horizontal="right")
    if last >= first + 1:
        tr = last + 1
        ws.cell(row=tr, column=1, value="Total")
        for col in range(2, len(heads) + 1):
            L = get_column_letter(col)
            ws.cell(row=tr, column=col, value=f"=SUM({L}{first + 1}:{L}{last})").number_format = BRL0
        for col in range(1, len(heads) + 1):
            c = ws.cell(row=tr, column=col)
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.font = Font(bold=True, color="FFFFFF", size=10)
    _widths(ws, [12] + [20] * len(used) + [16])
    _finish(ws, first, last, len(heads), "B")
    _ = tmap

    # ---- Pendências
    ws = wb.create_sheet("Pendências")
    heads = ["Loja", "Fornecedor / tipo", "Unidade consumidora", "Situação"]
    first = _title(ws, f"Pendências de {focus}", subtitle, len(heads))
    _header(ws, first, heads)
    row = first + 1
    for p in pending:
        _text(ws.cell(row=row, column=1), p["store"].code)
        _text(ws.cell(row=row, column=2), p["type"].name)
        _text(ws.cell(row=row, column=3), p["unit"].number if p.get("unit") else "")
        _text(ws.cell(row=row, column=4), p["text"])
        row += 1
    last = row - 1
    if last < first + 1:
        ws.cell(row=first + 1, column=1, value="Nenhuma pendência.").font = Font(color=GREY, italic=True)
        last = first + 1
    _body_style(ws, first + 1, last, len(heads))
    _widths(ws, [12, 36, 24, 52])
    _finish(ws, first, last, len(heads), "B", landscape=False)
    return _bytes(wb)


# --------------------------------------------------------------------------- lista de contas
def bills_workbook(rows: list[dict], filters_text: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Contas"
    heads = ["Mês de referência", "Loja", "Unidade consumidora", "Tipo", "Valor", "Vencimento", "Nota fiscal", "Origem"]
    first = _title(ws, "Contas e lançamentos · Economart", f"{filters_text} · {len(rows)} registro(s) · {_stamp()}", len(heads))
    _header(ws, first, heads)
    row = first + 1
    for r in rows:
        ref: date = r["reference"]
        ws.cell(row=row, column=1, value=date(ref.year, ref.month, 1)).number_format = "mm/yyyy"
        _text(ws.cell(row=row, column=2), r["store_code"])
        _text(ws.cell(row=row, column=3), r["unit_number"])
        _text(ws.cell(row=row, column=4), r["type_name"])
        ws.cell(row=row, column=5, value=float(r["value"]) if r["value"] is not None else None).number_format = BRL
        due = r["due_date"]
        ws.cell(row=row, column=6, value=datetime(due.year, due.month, due.day) if due else None).number_format = "dd/mm/yyyy"
        _text(ws.cell(row=row, column=7), r["invoice_number"])
        _text(ws.cell(row=row, column=8), "Foto/IA" if r["origin"] == "bill" else "Manual")
        row += 1
    last = row - 1
    if last < first + 1:
        ws.cell(row=first + 1, column=1, value="Nenhum registro com esses filtros.").font = Font(color=GREY, italic=True)
        last = first + 1
    _body_style(ws, first + 1, last, len(heads))
    for r_ in range(first + 1, last + 1):
        ws.cell(row=r_, column=1).alignment = Alignment(horizontal="center")
        ws.cell(row=r_, column=2).font = Font(bold=True, color=BLUE, size=10)
        ws.cell(row=r_, column=5).alignment = Alignment(horizontal="right")
    if last >= first + 1 and rows:
        tr = last + 1
        ws.cell(row=tr, column=1, value="Total")
        ws.cell(row=tr, column=5, value=f"=SUBTOTAL(109,E{first + 1}:E{last})").number_format = BRL
        for col in range(1, len(heads) + 1):
            c = ws.cell(row=tr, column=col)
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.font = Font(bold=True, color="FFFFFF", size=10)
    _widths(ws, [16, 12, 24, 30, 16, 14, 18, 11])
    _finish(ws, first, last, len(heads), "C")
    return _bytes(wb)
